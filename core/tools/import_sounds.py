"""Dev tool: writes the Index's sound data from the community's documents.

    python core/tools/import_sounds.py [labeling.xlsx] [names folder]

  labeling.xlsx  the HD2 modding community's "Helldivers 2 Archive Labeling" workbook (kept next to the project, in
                 _examples/). Its Audio tab names audio archives; its transcription tabs give the words of voice
                 lines by stream file id.
  names folder   filediver's hash lists (hashes.txt, cracked.txt), for the path names of the game's Wwise banks. The
                 game's own wwise_dep resources name the banks it has now (the index's "path").

It writes, in audiodiver/data:

  banks.json   the path names of the banks ("content/audio/wep_autocannon"). A bank's file id is the hash of its
               path, so a name finds its bank in any game version, whatever archives hold it.
  sounds.json  the Index's sound sets, by category and section. A set names its banks: by path name where that is
               known ("banks"), by file id where it isn't ("ids"). Both are the same in every game version; the
               archives that hold a bank are not, so they are never written down.
  voices.json  what each voice line says, by the id of its Sound object: the same object in every language, so the
               English words also label the line in the others.

The community names a thing by the archives it found it in, and an archive often holds more than its own bank (the
tutorial's has two dozen). So a row's bank is worked out: the bank all its archives have in common; where that leaves
several, the ones not known to be something else; and of those, the ones the fewest archives carry. BANK_OF settles
by hand the few rows this gets wrong. The tool prints every row it had to choose for.

The layout below (categories, sections, friendlier names) is curated here; the names of weapons, stratagems and
enemies are the community's. Needs the installed game, with the English voices, and openpyxl.
"""
import collections
import json
import re
import sys
from pathlib import Path

import openpyxl

from audiodiver import bank as wbank
from audiodiver import music
from audiodiver.game import murmur64_hash, stream_id
from audiodiver.index import AudioIndex, LANGUAGES
from audiodiver.library import slug
from audiodiver.titles import title

HERE = Path(__file__).parent
ROOT = HERE.parents[1]
DATA = HERE.parent / "audiodiver" / "data"
PREFIX = "content/audio/"
LANGUAGE_NAMES = {"Chinese": "cn", "English": "us", "French": "fr", "German": "de", "Italian": "it", "Japanese": "jp",
                  "Portuguese": "bp", "Spanish, Castilian": "es", "Spanish, Castillian": "es", "Spanish, LatAm": "ms"}
SPOKEN = re.compile(r"^(.*?) \((" + "|".join(re.escape(n) for n in LANGUAGE_NAMES) + r")\)$")

# the Audio tab's blocks of columns, left to right: where each one's rows go (category, section)
BLOCKS = ["weapons", "stratagems", "vehicles", None, None, "planets/environments", "planets/hazards",
          "planets/objectives", None, "enemies/terminids", "enemies/automatons", "enemies/illuminate", None]

CATEGORIES = [
    ("weapons", "Weapons", 60, "Everything a Helldiver carries and fires: primaries, sidearms, grenades and support weapons.", [
        ("weapons/primary", "Primary Weapons", ""),
        ("weapons/secondary", "Secondary Weapons", "Pistols, melee weapons and the other sidearms."),
        ("weapons/support", "Support Weapons", "The weapons a stratagem brings, and the ones found on the battlefield."),
        ("weapons/grenades", "Grenades", "Grenades, knives and the other throwables."),
        ("weapons/shared", "Shared", "")]),
    ("stratagems", "Stratagems", 250, "Orbitals, Eagles, sentries, emplacements, mines and backpacks.", [
        ("stratagems", "Stratagems", "")]),
    ("vehicles", "Vehicles", 215, "Vehicles and exosuits.", [
        ("vehicles", "Vehicles", "")]),
    ("helldivers", "Helldivers", 190, "The Helldiver's own voices, movement and gear.", [
        ("helldivers/voices", "Voices", "The four Helldiver voices, one part per installed language, and the grunts they share."),
        ("helldivers/body", "Movement and Body", "Footsteps on every surface, gear rustle and gore."),
        ("helldivers/gear", "Armor Passives and Boosters", "")]),
    ("allies", "Allies and Crew", 160, "Everyone on the Helldivers' side who talks: mission command, the Super Destroyer's crew, civilians and SEAF.", [
        ("allies/command", "Mission Command", "The voices on the radio during a mission."),
        ("allies/ship", "Super Destroyer", "The crew and the announcements on the ship."),
        ("allies/civilians", "Civilians and SEAF", ""),
        ("allies/training", "Training", "The training mission's voices and sounds.")]),
    ("enemies", "Enemies", 350, "The sounds of each enemy, by faction.", [
        ("enemies/terminids", "Terminids", ""),
        ("enemies/automatons", "Automatons", ""),
        ("enemies/illuminate", "Illuminate", "")]),
    ("planets", "Planets", 130, "Where missions happen: environments, hazards and objectives.", [
        ("planets/environments", "Environments", "Ambience and the things that stand around, by biome."),
        ("planets/hazards", "Hazards", ""),
        ("planets/objectives", "Objectives", "The machinery of each objective, by the faction whose missions have it.")]),
    ("interface", "Interface and Cutscenes", 320, "Menus, HUD sounds and cutscenes.", [
        ("interface/ui", "Interface", ""),
        ("interface/cutscenes", "Cutscenes", "")]),
]

# voices: (section, name, the bank's name in every language)
VOICES = [
    ("helldivers/voices", "Helldiver Voice 1", "helldiver_soldier_female1_VO"),
    ("helldivers/voices", "Helldiver Voice 2", "helldiver_soldier_male1_VO"),
    ("helldivers/voices", "Helldiver Voice 3", "helldiver_soldier_female2_VO"),
    ("helldivers/voices", "Helldiver Voice 4", "helldiver_purist_VO"),
    ("allies/command", "Democracy Officer (In Mission)", "Radio_Admiral_VO"),
    ("allies/command", "Mission Control", "Radio_Mission_Control_VO"),
    ("allies/command", "Eagle 1", "Radio_Eagle_Pilot_VO"),
    ("allies/command", "Extraction Pilot", "Radio_Extraction_Pilot_VO"),
    ("allies/command", "Shuttle Pilot", "Radio_Shuttle_Pilot_VO"),
    ("allies/command", "Objective Pilot", "Radio_Objective_Pilot_VO"),
    ("allies/command", "Objective Voices", "Objective_VO"),
    ("allies/ship", "Democracy Officer (On Ship)", "Shiphub_Admiral_VO"),
    ("allies/ship", "Ship Master", "Shiphub_Ship_Master_VO"),
    ("allies/ship", "Service Technician", "Shiphub_Service_Technician_VO"),
    ("allies/ship", "Hellbot", "Shiphub_Hellbot_VO"),
    ("allies/ship", "PA System and Ship TV", "Ship_Hub"),
    ("allies/civilians", "Civilians", "Mission_Civilian_VO"),
    ("allies/training", "General Brasch", "Tutorial_Instructor_VO"),
    ("allies/training", "Training Dummy", "Tutorial_Dummy_VO"),
    ("allies/training", "Training Extraction Pilot", "Tutorial_Extraction_Pilot_VO"),
    ("allies/training", "Training Facility Operator", "Tutorial_Facility_Operator_VO"),
    ("allies/training", "Helldiver Recruit", "Tutorial_Helldiver_Recruit_VO"),
    ("interface/cutscenes", "Cutscene Voices", "cutscenes"),
]

# banks placed by hand: path name (after content/audio/) -> (section, name). None: not a sound set (music, or empty)
PLACED = {
    "weapons_superearth": ("weapons", "Shared Weapon Sounds"),
    "stratagems_orbital_380mm_he_dss": ("stratagems", "Orbital 380mm HE Barrage (DSS)"),
    "vehicle_shared": ("vehicles", "Shared Vehicle Sounds"), "vehicle_shuttle": ("vehicles", "Shuttle"),
    "vehicle_oilrig": ("vehicles", "Oil Rig"), "vehicle_combat_walker": ("vehicles", "Exosuits"),
    "Helldiver_Standard_VO": ("helldivers/voices", "Helldiver Exertions"),
    "foley_player": ("helldivers/body", "Movement and Gear"), "gore_player": ("helldivers/body", "Gore"),
    "passive_armor_constitution": ("helldivers/gear", "Armor Passive: Constitution"),
    "passive_booster_hellpod_firebomb": ("helldivers/gear", "Booster: Hellpod Firebomb"),
    "passive_booster_hellpod_smokepods": ("helldivers/gear", "Booster: Hellpod Smoke Pods"),
    "passive_booster_hellpod_stun": ("helldivers/gear", "Booster: Hellpod Stun"),
    "vo_speakers": ("allies/command", "Loudspeakers"), "vo_radio": ("allies/command", "Radio"),
    "vo_objects": ("allies/command", "Talking Objects"),
    "Mission_Civilian_Exertions": ("allies/civilians", "Civilian Exertions"),
    "foley_npcs": ("allies/civilians", "NPC Movement"), "foley_seafs": ("allies/civilians", "SEAF Movement"),
    "tutorial": ("allies/training", "Training Sounds"), "tutorial_outro": ("allies/training", "Training Outro"),
    "foley_bugs": ("enemies/terminids", "Terminid Movement"), "gore_bugs": ("enemies/terminids", "Terminid Gore"),
    "vo_bugs": ("enemies/terminids", "Terminid Voices"),
    "foley_bots": ("enemies/automatons", "Automaton Movement"), "gore_bots": ("enemies/automatons", "Automaton Gore"),
    "weapons_bots": ("enemies/automatons", "Automaton Weapons"),
    "bots_colony_spawner": ("enemies/automatons", "Colony Spawner"),
    "bots_fabricator": ("enemies/automatons", "Fabricator"),
    "cyborg_elite": ("enemies/automatons", "Cyborg Elite"), "cyborg_siegeengine": ("enemies/automatons", "Vox Engine"),
    "foley_illuminate": ("enemies/illuminate", "Illuminate Movement"), "gore_illuminate": ("enemies/illuminate", "Illuminate Gore"),
    "weapons_illuminate": ("enemies/illuminate", "Illuminate Weapons"),
    "dropship_illuminate": ("enemies/illuminate", "Illuminate Dropship"),
    "vo_illuminate": ("enemies/illuminate", "Illuminate Voices"),
    "illuminates_bladed": ("enemies/illuminate", "Wretch"), "illuminates_helmetguy": ("enemies/illuminate", "Crusher"),
    "illuminates_exomech": ("enemies/illuminate", "Gatekeeper, Veracitor"),
    "destruction": ("planets/environments", "Destruction"),
    "env_bots": ("planets/environments", "Automaton Worlds"), "env_bugs": ("planets/environments", "Terminid Worlds"),
    "env_illuminate": ("planets/environments", "Illuminate Worlds"), "env_ship": ("planets/environments", "Super Destroyer"),
    "env_shared": ("planets/environments", "Shared Environment Sounds"), "env_sandy": ("planets/environments", "Desert"),
    "obj_hellbomb": ("planets/objectives", "Hellbomb"), "obj_samples": ("planets/objectives", "Samples"),
    "obj_generic_terminal": ("planets/objectives", "Terminal"), "obj_colony_relay": ("planets/objectives", "Colony Relay"),
    "ui_mission": ("interface/ui", "Mission Interface"), "ui_ship": ("interface/ui", "Ship Interface"),
    "game_init": ("interface/ui", "Game Init"), "intro_logos": ("interface/ui", "Intro Logos"),
    "rev_gameplay": ("interface/ui", "Reverb: Missions"), "rev_ship_hub": ("interface/ui", "Reverb: Ship"),
    "Helldiver_IR": ("interface/ui", "Reverb: Helmet"),
    "cutscenes_sfx": ("interface/cutscenes", "Cutscene Sound Effects"),
    "Init": None, "music_init": None, "music_ship": None, "music_mission_bot": None, "music_mission_bot_infiltration": None,
    "music_mission_bug": None, "music_mission_bug_hiveworld": None, "music_mission_illuminate": None,
    "music_mission_tutorial": None,
}
# names a set had before (renamed by the user): written as its "was", so a project that listed it by the old name
# finds it under the new one (library.py, Library.renamed)
WAS = {"Vox Engine": ["Cyborg Siege Engine"], "Game Init": ["Game Start"]}
# named rows of the Audio tab placed by hand: the community's name -> (section, name)
RENAMED = {
    "SEAF Exertions": ("allies/civilians", "SEAF Exertions"),
    "Experimental Infusion": ("helldivers/gear", "Booster: Experimental Infusion"),
    "Stratagem Shared Sounds": ("stratagems", "Shared Stratagem Sounds"),
}
# rows left out: their banks are placed by name above (Helldiver_Standard_VO, Mission_Civilian_Exertions), and the
# archives the community lists for them hold a good deal else
DROPPED = {"Helldiver Exertions", "Civilian Exertions"}
# rows whose archive holds several banks, settled by hand: the community's name -> the bank's path name
BANK_OF = {
    "LAS-98 Laser Cannon": "wep_laser_cannon", "LAS-99 Quasar Cannon": "wep_laser_pulse_cannon",
    "LAS-7 Dagger": "las_7", "PLAS-1 Scorcher": "wep_plas1_single", "R-63 Diligence": "wep_r63_vigilance",
    "EAT-17 Expendable Anti-tank": "wep_eat17_expendable_at_rifle", "Dropship": "bots_dropship",
    "Scout Strider": "bots_scoutwalker",
}
WORDS = {"aa": "AA", "sam": "SAM", "samsite": "SAM Site", "icbm": "ICBM", "seaf": "SEAF", "tcs": "TCS", "harddrive": "Hard Drive",
         "foggenerator": "Fog Generator", "bugcluster": "Bug Cluster", "dustdevil": "Dust Devil", "firetornado": "Fire Tornado",
         "explosivemushroom": "Explosive Mushroom", "frozenflowers": "Frozen Flowers", "meteorstorm": "Meteor Storm",
         "toxchimney": "Toxic Chimney", "volcanicactivity": "Volcanic Activity", "hiveworld": "Hive World",
         "jumppack": "Jump Pack", "inf": "Infiltration", "il": "Illuminate", "cy": "Automatons", "bug": "Terminids",
         "bot": "Automatons", "gen": "", "mud": "Mud", "ice": "Ice", "dropoff": "Drop Off", "shutdown": "Shut Down",
         "gre": "Grenade", "smg": "SMG", "frv": "FRV", "sg": "Shotgun", "ar": "AR", "at": "AT"}
# Each weapon set's slot in the game's loadout and its in-game name: the name the Audio tab (or the bank's path name)
# gives it -> (section, name shown). The names come from the workbook's "Weapons & Attachment IDs" tab (the weapon
# whose archive holds the bank) and the Helldivers 2 wiki's weapon lists, which also give the slots; the Scope of
# Authority weapons (R-40, AMW-20, LAS-44) were named by the user. The LAS-99 Quasar Cannon is wep_laser_pulse_cannon,
# in the archive the tab calls "LAS-98 Laser Cannon & LAS-99 Quasar Cannon"; wep_laser_charge_rifle, which sounds alike,
# is in the LAS-22 Shear's. Weapons that aren't a Helldiver's go where their users are (exosuits and FRVs to Vehicles,
# the SEAF's to Allies). The tool stops on a weapon set this table doesn't place (a new one after a game update): look
# it up and add it here.
P, S, U, G = "weapons/primary", "weapons/secondary", "weapons/support", "weapons/grenades"
WEAPONS = {
    "AC-8 Autocannon": (U, None), "APW-1 Anti-Materiel Rifle": (U, None), "AR Risk": (P, "AR-11 Arbitrator"),
    "AR-2 Coyote": (P, None), "AR-23 Liberator": (P, None), "AR-23A Liberator Carbine": (P, None), "AR-32 Pacifier": (P, None),
    "AR-59 Suppressor": (P, None), "AR-61 Tenderizer": (P, None), "AR/GL-21 One Two": (P, "AR/GL-21 One-Two"),
    "ARC-12 Blitzer": (P, None), "ARC-3 Arc Thrower": (U, None),
    "Belt-Fed Grenade Launcher": (U, "GL-28 Belt-Fed Grenade Launcher"), "BR-14 Adjudicator": (P, None),
    "Breaching Hammer": (U, "CQC-20 Breaching Hammer"), "Break-Action Shotgun": (U, "SG-88 Break-Action Shotgun"),
    "Broomhandle": (S, "P-69 Veto"), "Bull Machine Pistol": (S, None), "CB-9 Exploding Crossbow": (P, None),
    "Combat Walker Breacher Flak Cannon": ("vehicles", "EXO-55 Breakthrough Flak Cannon"),
    "Combat Walker Flamethrower": ("vehicles", "EXO-51 Lumberer Flamethrower"),
    "Combat Walker Lumberer Cannon": ("vehicles", "EXO-51 Lumberer Cannon"),
    "CQC-19 Stun Lance": (S, None), "CQC-2 Saber": (S, None), "CQC-30 Stun Baton": (S, None), "CQC-42 Machete": (S, None),
    "CQC-5 Combat Hatchet": (S, None), "CQC-9 Defoliation Tool": (U, None), "Double Freedom": (P, "DBS-2 Double Freedom"),
    "Drake Rifle": (P, "R-4 Hyena"),  # its archive is the R-63 Diligence's, with the R-63CS and the R-4 Hyena
    "EAT-17 Expendable Anti-Tank": (U, None), "EAT-700 Expendable Napalm": (U, None),
    "Entrenching Tool": (U, "CQC-72 Entrenchment Tool"), "Expendable Machinegun": (U, "MGX-42 Bullet Storm"),
    "FAF-14 Spear": (U, None), "FLAM-40 Flamethrower": (U, None), "FLAM-66 Torcher": (P, None),
    "FRV Heavy Flamer": ("vehicles", "M-104 Incinerator FRV Flamer"),
    "FRV Supply Autoturret": ("vehicles", "M-103 Supply FRV Autoturret"),
    "G-109 Urchin": (G, None), "G-12 High Explosive, G-16 Impact Grenades": (G, None), "G-123 Thermite Grenade": (G, None),
    "G-13 Incendiary Impact, G-10 Incendiary Grenades": (G, None), "G-142 Pyrotech Grenade": (G, None),
    "G-23 Stun Grenade": (G, None), "G-3 Smoke Grenade": (G, None), "G-31 Arc Grenade": (G, None), "G-4 Gas Grenade": (G, None),
    "G-50 Seeker": (G, None), "G-6 Frag Grenade": (G, None), "G-7 Pineapple": (G, None), "G/SH-39 Shield Grenade": (G, None),
    "GL-21 Grenade Launcher": (U, None), "GL-52 De-Escalator": (U, None), "GP-20 Ultimatum": (S, None),
    "GP-31 Grenade Pistol": (S, None), "GR-8 Recoilless Rifle": (U, None),
    "Grenade AT Explosive": (G, "G-48 Giga Grenade"), "Grenade AT Self Destruct Drone": (G, "G-60 Anti-Tank Seeker"),
    "Grenade Caltrops": (G, "G-11 Caltrops"), "Grenade Melta Mine": (G, "G/40-K Melta Mine"),
    "Grenade Molotov": (G, "G-8 Immolation"), "Heavy Flamer": (U, "B/FLAM-80 Cremator"),
    "Hornet Pistol": (S, "P-34 Breacher"), "Hotshot Marksman Rifle": (P, "R/40-K Hot-Shot Marksman Rifle"),
    "JAR-5 Dominator": (P, None), "K-2 Throwing Knife": (G, None), "LAS-16 Sickle": (P, None),
    "LAS-16 Trident": (P, "LAS-13 Trident"), "LAS-17 Double-Edge Sickle": (P, None), "LAS-5 Scythe": (P, None),
    "LAS-58 Talon": (S, None), "LAS-7 Dagger": (S, None), "LAS-98 Laser Cannon": (U, None), "LAS-99 Quasar Cannon": (U, None),
    "Laser Charge Rifle": (P, "LAS-22 Shear"), "Laser Rifle Ion": (U, "LAS-44 Nova"), "Leveller": (U, "EAT-411 Leveller"),
    "Liberator Legacy": (P, None), "M-1000 Maxigun": (U, None), "M-105 Stalwart": (U, None), "M6C/SOCOM Pistol": (S, None),
    "M7S SMG": (P, None), "M90A Shotgun": (P, None), "MA5C Assault Rifle": (P, None), "Meltagun": (U, "40-K Meltagun"),
    "MG-206 Heavy Machine Gun": (U, None), "MG-43 Machine Gun": (U, None), "MLS-4X Commando": (U, None),
    "MP-98 Knight": (P, None), "Ombudsman": (S, "P-41 Ombudsman"), "One True Flag": (U, "CQC-1 One True Flag"),
    "P-11 Stim Pistol": (S, None), "P-113 Verdict": (S, None), "P-19 Redeemer": (S, None), "P-2 Peacemaker": (S, None),
    "P-35 Re-Educator": (S, None), "P-72 Crisper": (S, None), "P-8 Senator": (S, "P-4 Senator"),
    "Pistol Shark": (S, "P/40-K Bolt Pistol"), "PLAS-1 Scorcher": (P, None), "PLAS-101 Purifier": (P, None),
    "PLAS-15 Loyalist": (S, None), "PLAS-39 Accelerator Rifle": (P, None), "PLAS-45 Epoch": (U, None),
    "PT-92 Warrant": (S, "P-92 Warrant"), "Pump Grenade Launcher": (P, "GL-15 Evictor"), "R-2 Amendment": (P, None),
    "R-2124 Constitution": (P, None), "R-36 Eruptor": (P, None), "R-6 Deadeye": (P, None), "R-63 Diligence": (P, None),
    "R-72 Censor": (P, None), "RL-77 Airburst Rocket Launcher": (U, None), "RS-422 Railgun": (U, None),
    "S-11 Speargun": (U, None),
    "SEAF Defender": ("allies/civilians", None), "SEAF Expendable Anti-Tank": ("allies/civilians", None),
    "SEAF Grenade": ("allies/civilians", None), "SEAF Liberator": ("allies/civilians", None),
    "SEAF Machinegun": ("allies/civilians", None), "SEAF Stim Pistol": ("allies/civilians", None),
    "SG-20 Halt": (P, None), "SG-22 Bushwhacker": (S, None), "SG-255 Breaker": (P, "SG-225 Breaker"),
    "SG-451 Cookout": (P, None), "SG-8 Punisher": (P, None), "Shared Weapon Sounds": ("weapons/shared", None),
    "Shotgun Trenchgun": (P, "SG-97 Sweeper"), "Smart Pistol Missile": (S, "P-33 Missile Pistol"),
    "SMG Flamer": (P, "SMG/FLAM-34 Stoker"), "SMG Laser Blaster": (P, "LAS-12 Sai"), "SMG Solvent": (P, "SMG-203 Gallant"),
    "SMG-32 Reprimand": (P, None), "SMG-37 Defender": (P, None), "SMG-72 Pummeler": (P, None),
    "Smokescreen": (G, "G-89 Smokescreen"), "Sniper Rifle Duo": (U, "AMW-20 Apex"), "Sniper Rifle Simo": (P, "R-40 Sightliner"),
    "StA-11 SMG": (P, None), "StA-52 Assault Rifle": (P, None), "StA-X3 W.A.S.P. Launcher (WASP Launcher)": (U, None),
    "TED-63 Dynamite Grenade": (G, None), "TM-1 Lure Mine": (G, None), "TX-41 Sterilizer": (U, None), "VG-70 Variable": (P, None),
}


def slot_weapons(sections: dict):
    """Moves the weapon sets out of the one "weapons" section the Audio tab fills, into their loadout slots (WEAPONS),
    under their in-game names. A set renamed or moved keeps its old name and id as its "was", so a project that
    listed it finds it (library.py)."""
    unplaced = sorted(set(sections.get("weapons", {})) - set(WEAPONS))
    assert not unplaced, f"weapon sets WEAPONS doesn't place (look them up and add them): {unplaced}"
    for name, e in sections.pop("weapons", {}).items():
        section, shown = WEAPONS[name]
        shown = shown or name
        was = [name] if shown != name else []
        if not section.startswith("weapons/"):
            was = [f"set:weapons/{slug(name)}"]  # another category: by the id it had in Weapons
        moved = {**e, "name": shown, **({"was": sorted(set(e.get("was", [])) | set(was))} if was else {})}
        assert shown not in sections[section], f"two weapon sets named {shown} in {section}"
        sections[section][shown] = moved


# weapon banks without the wep_ prefix (the game's own names, from its wwise_dep resources)
SIDEARMS = {"broomhandle", "bull_machine_pistol", "hornet_pistol", "ombudsman", "pistol_shark", "smart_pistol_missile"}
FACTIONS = {"bug": "Terminids", "cy": "Automatons", "bot": "Automatons", "il": "Illuminate"}


def words(tail: str) -> str:
    return title(" ".join(WORDS.get(w, w) for w in tail.split("_") if WORDS.get(w, w)))


def pretty(tail: str) -> tuple[str, str] | None:
    """(section, name) for a bank's path name, by its prefix; None for a bank that isn't a sound set."""
    if tail in PLACED:
        return PLACED[tail]
    kind, _, rest = tail.partition("_")
    if kind == "env":
        return "planets/environments", words(rest)
    if kind == "haz":
        return "planets/hazards", words(rest)
    if kind == "obj":
        faction, _, what = rest.partition("_")
        if faction in FACTIONS or faction == "gen":
            return "planets/objectives", (f"{FACTIONS[faction]}: " if faction in FACTIONS else "") + words(what)
        return "planets/objectives", words(rest)
    if tail.startswith("foley_player_surface_media_"):  # "concrete_stone": Concrete and Stone
        return "helldivers/body", "Footsteps: " + words(tail.removeprefix("foley_player_surface_media_")).replace(" ", " and ", 1)
    # banks no community list names (new in an update, most often): placed and named by their path name alone
    if kind == "wep" or tail in SIDEARMS:
        return "weapons", words(rest if kind == "wep" else tail)
    if tail.startswith("seaf_weapons_"):
        return "weapons", "SEAF " + words(tail.removeprefix("seaf_weapons_"))
    if kind == "stratagems":
        return "stratagems", words(rest)
    if kind == "vehicle":
        return "vehicles", words(rest)
    if kind == "bugs":
        return "enemies/terminids", words(rest)
    if kind == "bots":
        return "enemies/automatons", words(rest)
    return None


def audio_tab(path) -> list[list[tuple[str, list[str]]]]:
    rows = [list(r) for r in openpyxl.load_workbook(path, read_only=True, data_only=True)["Audio"].iter_rows(values_only=True)]
    starts = [i for i, c in enumerate(rows[0]) if c == "Name"]
    blocks = []
    for n, s in enumerate(starts):
        e = starts[n + 1] if n + 1 < len(starts) else len(rows[0])
        out = []
        for r in rows[1:]:
            cells = [("" if c is None else str(c).strip()) for c in r[s:e]]
            if cells and cells[0]:
                out.append((cells[0], [c for c in cells[1:] if re.fullmatch(r"[0-9a-f]{16}", c)]))
        blocks.append(out)
    return blocks


def transcriptions(path) -> dict[int, str]:
    """{stream file id: what the line says}, from the workbook's two transcription tabs."""
    wb, text = openpyxl.load_workbook(path, read_only=True, data_only=True), {}

    def take(fid, words_):
        try:
            fid = int(str(fid).strip().split(".")[0])
        except (TypeError, ValueError):
            return
        words_ = " ".join(str(words_).split()) if words_ is not None else ""
        if fid > 2 ** 32 and words_ and words_.lower() not in ("none", "?", "-"):
            text.setdefault(fid, words_)
    rows = [list(r) for r in wb["Audio Transcriptions"].iter_rows(values_only=True)]
    cols = [i for i, c in enumerate(rows[0]) if c == "File ID"]
    for r in rows[1:]:
        for c in cols:
            if c + 1 < len(r):
                take(r[c], r[c + 1])
    for r in wb["transcriptionHub"].iter_rows(values_only=True, min_row=2):
        if len(r) >= 2:
            take(r[0], r[1])
    return text


# bank names no list has, found by trying names against the game's bank ids (murmur64): the newer Illuminate units'
# banks use a plural prefix. The workbook's "Enemies & NPC IDs" tab names their archives.
KNOWN = [PREFIX + t for t in ("illuminates_bladed", "illuminates_exomech", "illuminates_helmetguy")]


def bank_names(folder, blocks, ix) -> list[str]:
    names = {n.split(" ")[0] for b in blocks for n, _ in b if n.startswith(PREFIX)} | set(KNOWN)
    names |= {b["path"] for b in ix.banks.values() if b.get("path")}  # as the game spells them (index.py)
    for f in ("hashes.txt", "cracked.txt"):
        p = Path(folder, f)
        if p.exists():
            for line in p.read_text(encoding="utf-8", errors="replace").splitlines():
                line = line.strip()
                if line.startswith(PREFIX) and not re.fullmatch(PREFIX + r"(\w+/)?\d+", line):
                    names.add(line)
    return sorted(names, key=str.lower)


def borrowed_calls(ix: AudioIndex, said: dict[int, str]) -> dict[int, str]:
    """{Sound id: a line's call} for the voices nobody transcribed (Helldiver Voice 1), from a voice of the same
    section that is built the same way. The Helldiver voices are: their Sounds come in the same order, and the n-th
    Sound of each sits in the n-th group of takes, the same call. The takes of a call are worded and ordered
    differently in each voice, so this names the call and nothing more."""
    def layout(tail):
        bid = murmur64_hash(f"{PREFIX}us/{tail}".encode())
        if bid not in ix.banks:
            return None
        out = []  # (Sound id, its parent: the group of takes it is in)
        for t, oid, body in music.iter_hirc(wbank.hirc(ix.bank_files(bid)[0])):
            if t == wbank.SOUND:
                try:
                    out.append((oid, music.read_node_base_params(music.Reader(body, 18))[1]))
                except Exception:
                    out.append((oid, 0))
        return out

    layouts = {tail: (section, lay) for section, _, tail in VOICES if (lay := layout(tail))}
    calls = {}
    for tail, (section, lay) in layouts.items():
        if any(oid in said for oid, _ in lay):
            continue  # transcribed
        for other, (o_section, o_lay) in layouts.items():
            if o_section != section or len(o_lay) != len(lay) or not any(oid in said for oid, _ in o_lay):
                continue
            pairs = {}  # the same layout: each group of this voice lines up with one group of the other
            if any(pairs.setdefault(p, q) != q for (_, p), (_, q) in zip(lay, o_lay)):
                continue
            before = len(calls)
            for (oid, _), (o_oid, _) in zip(lay, o_lay):
                if oid not in calls and o_oid in said:
                    calls[oid] = said[o_oid]
            if len(calls) > before:
                print(f"  {tail}: {len(calls) - before} of {len(lay)} lines named after {other}")
    return calls


def main():
    sheet = sys.argv[1] if len(sys.argv) > 1 else str(ROOT.parent / "_examples" / "Helldivers 2 Archive Labeling.xlsx")
    names_dir = sys.argv[2] if len(sys.argv) > 2 else str(ROOT.parents[1] / "Script" / "_shared" / "filediver" / "hashes")
    blocks = audio_tab(sheet)
    assert len(blocks) == len(BLOCKS), f"the Audio tab has {len(blocks)} blocks of columns now, not {len(BLOCKS)}: check BLOCKS"
    ix = AudioIndex(log=print)
    names = bank_names(names_dir, blocks, ix)
    # a name is a bank's as it is, or under a language ("content/audio/us/cutscenes"); keep the plain form of both
    tails = sorted({n.removeprefix(PREFIX).split("/")[-1] for n in names if n.count("/") in (2, 3)}, key=str.lower)
    by_id = {murmur64_hash((PREFIX + t).encode()): t for t in tails}
    voice_ids = {murmur64_hash(f"{PREFIX}{code}/{t}".encode()) for _, _, t in VOICES for code in LANGUAGES}
    arch = collections.defaultdict(set)
    for bid, b in ix.banks.items():
        for a in b["archives"]:
            arch[a].add(bid)
    has_audio = lambda bid: bid in ix.banks and bool(ix.banks[bid]["refs"] or ix.banks[bid]["media"])
    in_game = lambda tail: has_audio(murmur64_hash((PREFIX + tail).encode()))

    sections = collections.defaultdict(dict)   # section id -> {name: entry}
    claimed = set()                            # bank tails placed in a set

    def add(section, name, **entry):
        e = sections[section].setdefault(name, {"name": name})
        for k, v in entry.items():
            e[k] = sorted(set(e.get(k, [])) | set(v)) if isinstance(v, list) else v

    for section, name, tail in VOICES:
        claimed.add(tail)
        if any(has_audio(murmur64_hash(f"{PREFIX}{code}/{tail}".encode())) for code in LANGUAGES):
            add(section, name, voice=tail)
    seaf = collections.defaultdict(list)       # the one voice whose bank name isn't known: by the archives it is in
    rows = []                                  # (the community's name, section, name shown, archives): things named by archive
    for block, where in zip(blocks, BLOCKS):
        for name, archives in block:
            if name.startswith(PREFIX):  # a bank's path name: the community lists it, so it stays even when this
                tail = name.split(" ")[0].removeprefix(PREFIX).split("/")[-1]  # game version doesn't have it
                place = pretty(tail)
                if place:
                    add(place[0], place[1], banks=[PREFIX + tail], **({"was": WAS[place[1]]} if place[1] in WAS else {}))
                    claimed.add(tail)
                continue
            m = SPOKEN.match(name)
            if m:  # a voice in one language: the VOICES above cover them by bank name, except the SEAF soldiers
                if m.group(1).startswith("SEAF Soldier"):
                    seaf[LANGUAGE_NAMES[m.group(2)]] += archives
                continue
            if name in DROPPED:
                continue
            section, shown = RENAMED.get(name, (where, title(name)))
            if not section:
                print(f"  not placed: {name} {archives}")
                continue
            rows.append((name, section, shown, archives))
    assert not set(BANK_OF) - {r[0] for r in rows}, f"BANK_OF names rows the Audio tab no longer has: {set(BANK_OF) - {r[0] for r in rows}}"

    # --- the bank of each thing named by its archives
    def held(archives):
        """(how many of these archives hold each bank, the banks they have in common): in all of them, else (a
        thing whose variants each bring an archive of their own) in at least half."""
        here = [a for a in archives if a in arch]
        count = collections.Counter(b for a in here for b in arch[a] if has_audio(b))
        return count, [b for b, n in count.items() if n == len(here)] or [b for b, n in count.items() if n * 2 >= len(here)]

    # known to be something else: banks placed by hand, the voices and the music; then what step 1 gives to a row
    owned = {b for b, t in by_id.items() if t in PLACED} | voice_ids | set(ix.music_banks)
    chosen = {}
    for i, (name, _section, _shown, archives) in enumerate(rows):  # 1. settled by hand, or one bank in common
        if name in BANK_OF:
            assert in_game(BANK_OF[name]), f"BANK_OF: {BANK_OF[name]} is not in this game version"
            chosen[i] = [murmur64_hash((PREFIX + BANK_OF[name]).encode())]
        else:
            common = held(archives)[1]
            if len(common) == 1 and common[0] not in owned:
                chosen[i] = common
    owned |= {b for bs in chosen.values() for b in bs}
    in_rows = collections.Counter(b for r in rows for b in {x for a in r[3] for x in arch.get(a, ())})
    print("rows whose archives hold several banks:")
    for i, (name, _section, _shown, archives) in enumerate(rows):  # 2. what is left once those are taken out
        if i in chosen:
            continue
        count, common = held(archives)
        free = [b for b in common if b not in owned] or [b for b in count if b not in owned]
        if not free:
            print(f"  not placed (no bank of its own in this game version): {name} {archives}")
            continue
        rank = lambda b: (len(ix.banks[b]["archives"]), in_rows[b])  # in the fewest archives, then named by the fewest rows
        chosen[i] = sorted(b for b in free if rank(b) == min(map(rank, free)))
        print(f"  {name}: {', '.join(by_id.get(b, str(b)) for b in chosen[i])}"
              f"  (of {', '.join(by_id.get(b, str(b)) for b in sorted(count, key=lambda b: (-count[b], b))[:8])})")
    merged = {}  # (section, banks) -> its rows: two names for one bank are one set ("Rocket Strider, Scout Strider")
    for i, bs in chosen.items():
        merged.setdefault((rows[i][1], tuple(bs)), []).append(rows[i][2])
    for (section, bs), names_ in merged.items():
        named = sorted(by_id[b] for b in bs if b in by_id)
        add(section, ", ".join(names_), **({"banks": [PREFIX + t for t in named]} if named else {}),
            **({"ids": [str(b) for b in bs if b not in by_id]} if len(named) < len(bs) else {}))
        claimed.update(named)
    if seaf:
        add("allies/civilians", "SEAF Soldier", archives_by_language={k: sorted(set(v)) for k, v in sorted(seaf.items())})
    for tail in tails:  # banks nobody listed: placed by their own name, if this game version has them
        if tail not in claimed and in_game(tail) and (place := pretty(tail)):
            add(place[0], place[1], banks=[PREFIX + tail])
            claimed.add(tail)
    left = [t for t in tails if t not in claimed and t not in PLACED and in_game(t)]
    slot_weapons(sections)

    out = {"about": "The Index's sound sets (library.py), written by tools/import_sounds.py from the HD2 modding community's "
                    "archive labeling workbook and filediver's name lists. A set is the banks it names: 'banks' by path "
                    "name, 'ids' by file id where the path name isn't known, 'voice' one bank per language "
                    "(content/audio/<language>/<name>). The one voice without a known name lists the archives the "
                    "community found it in, per language ('archives_by_language'). 'hue' colours a category's tags.",
           "categories": []}
    for cid, cname, hue, about, secs in CATEGORIES:
        cat = {"id": cid, "name": cname, "hue": hue, "about": about, "sections": []}
        for sid, sname, sabout in secs:
            entries = sorted(sections.pop(sid, {}).values(), key=lambda e: e["name"].lower())
            cat["sections"].append({"id": sid, "name": sname, "about": sabout, "entries": entries})
        out["categories"].append(cat)
    assert not sections, f"sections without a category: {list(sections)}"

    lines = ["{", f' "about": {json.dumps(out["about"])},', ' "categories": [']
    for ci, cat in enumerate(out["categories"]):
        lines.append("  {" + ", ".join(f'"{k}": {json.dumps(cat[k], ensure_ascii=False)}' for k in ("id", "name", "hue", "about")) + ', "sections": [')
        for si, sec in enumerate(cat["sections"]):
            lines.append("   {" + ", ".join(f'"{k}": {json.dumps(sec[k], ensure_ascii=False)}' for k in ("id", "name", "about")) + ', "entries": [')
            lines += [f"    {json.dumps(e, ensure_ascii=False)}{',' if n + 1 < len(sec['entries']) else ''}" for n, e in enumerate(sec["entries"])]
            lines.append("   ]}" + ("," if si + 1 < len(cat["sections"]) else ""))
        lines.append("  ]}" + ("," if ci + 1 < len(out["categories"]) else ""))
    lines += [" ]", "}"]
    (DATA / "sounds.json").write_text("\n".join(lines) + "\n", encoding="utf-8", newline="\n")
    (DATA / "banks.json").write_text(json.dumps({
        "about": "Path names of the game's Wwise banks, from filediver's name lists and the community's archive labeling "
                 "workbook. A bank's file id is the hash of its path; voice banks sit under a language "
                 "(content/audio/us/<name>).", "names": tails}, indent=0) + "\n", encoding="utf-8", newline="\n")

    # --- voice lines: the workbook's words by stream file id -> by Sound object id, through the English banks
    text, said = transcriptions(sheet), {}
    for bid, b in ix.banks.items():
        if b.get("lang") != "us":
            continue
        for t, oid, body in music.iter_hirc(wbank.hirc(ix.bank_files(bid)[0])):
            if t == wbank.SOUND:
                line = text.get(stream_id(wbank.SOURCE.unpack_from(body, 0)[2], "us"))
                if line:
                    said[oid] = line
    calls = borrowed_calls(ix, said)
    (DATA / "voices.json").write_text(json.dumps({
        "about": "What voice lines say, by the id of their English Sound object (the other languages map to it by place), from the "
                 "transcriptions in the community's archive labeling workbook. Calls: for a voice nobody transcribed, "
                 "what a sibling voice built the same way says at the same place. They name the line's call, not its "
                 "words.",
        "lines": {str(k): v for k, v in sorted(said.items())},
        "calls": {str(k): v for k, v in sorted(calls.items())}}, ensure_ascii=False, indent=0) + "\n", encoding="utf-8", newline="\n")

    n_sets = sum(len(s["entries"]) for c in out["categories"] for s in c["sections"])
    print(f"{n_sets} sound sets, {len(tails)} bank names, {len(said)} of {len(text)} voice lines matched to a Sound, "
          f"{len(calls)} named by a sibling voice")
    print("named banks left out of every set:", ", ".join(left) or "none")
    for c in out["categories"]:
        print(f"  {c['name']}: " + "; ".join(f"{s['name']} {len(s['entries'])}" for s in c["sections"]))


if __name__ == "__main__":
    main()
